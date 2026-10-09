"""Read the passenger list from an Excel file (.xlsx). Columns found by their heading, in any order."""
import io
from dataclasses import dataclass, field

from openpyxl import load_workbook

MAX_ROWS = 5000

# heading (lower case, spaces and dashes as underscores) -> field
HEADINGS = {
    "employee_id": "employee_id", "emp_id": "employee_id", "employee_no": "employee_id", "id": "employee_id",
    "employee": "employee_id", "staff_id": "employee_id",
    "name": "name", "employee_name": "name", "full_name": "name",
    "department": "department", "dept": "department",
    "phone": "phone", "mobile": "phone", "phone_no": "phone", "phone_number": "phone", "mobile_no": "phone",
}


class ImportFileError(ValueError):
    """The file cannot be used at all (wrong type, no headings, too big)."""


@dataclass
class Row:
    line: int
    employee_id: str
    name: str
    department: str | None
    phone: str | None


@dataclass
class Parsed:
    rows: list[Row] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # Excel stores 2210 as 2210.0
        v = int(v)
    return str(v).strip()


def parse_passenger_xlsx(data: bytes) -> Parsed:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("This is not a readable .xlsx file. Save the sheet as Excel Workbook (.xlsx).")
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        head = next(rows)
    except StopIteration:
        raise ImportFileError("The file is empty.")
    cols: dict[str, int] = {}
    for i, h in enumerate(head):
        key = HEADINGS.get(_text(h).lower().replace(" ", "_").replace("-", "_"))
        if key and key not in cols:
            cols[key] = i
    missing = [f for f in ("employee_id", "name") if f not in cols]
    if missing:
        raise ImportFileError("The first row must have the headings Employee ID and Name (Department and Phone are optional).")

    out = Parsed()
    seen: dict[str, int] = {}
    for n, raw in enumerate(rows, start=2):
        if n - 1 > MAX_ROWS:
            raise ImportFileError(f"The file has more than {MAX_ROWS} rows.")
        cell = lambda f: _text(raw[cols[f]]) if f in cols and cols[f] < len(raw) else ""
        emp, name, dept, phone = cell("employee_id"), cell("name"), cell("department"), cell("phone")
        if not (emp or name or dept or phone):
            continue  # empty line
        problem = None
        if not emp:
            problem = "Employee ID is empty"
        elif not name:
            problem = "Name is empty"
        elif len(emp) > 40:
            problem = "Employee ID is longer than 40 characters"
        elif len(name) > 120:
            problem = "Name is longer than 120 characters"
        elif len(dept) > 120:
            problem = "Department is longer than 120 characters"
        elif len(phone) > 30:
            problem = "Phone is longer than 30 characters"
        elif emp.lower() in seen:
            problem = f"Employee ID {emp} is already in row {seen[emp.lower()]}"
        if problem:
            out.errors.append({"row": n, "employee_id": emp or None, "problem": problem})
            continue
        seen[emp.lower()] = n
        out.rows.append(Row(n, emp, name, dept or None, phone or None))
    return out
