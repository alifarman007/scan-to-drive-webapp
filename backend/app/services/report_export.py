"""Excel (openpyxl) and PDF (ReportLab) files for a report."""
import io
import re
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import settings
from app.services.reports import Report
from app.services.timeutil import LOCAL_TZ

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_TYPE = "application/pdf"

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(v):
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def _sub_title(r: Report) -> str:
    parts = [f"Period: {r.date_from.isoformat()} to {r.date_to.isoformat()} (Bangladesh time)"]
    parts += [f"{k}: {v}" for k, v in r.filters.items()]
    return "   |   ".join(parts)


def _made_at() -> str:
    return "Made " + datetime.now(timezone.utc).astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M")


def _sheet_name(title: str | None, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", " ", title or "Report")[:31].strip() or "Report"
    name, i = base, 2
    while name in used:
        name = f"{base[:28]} {i}"
        i += 1
    used.add(name)
    return name


def _put(ws, row: int, col: int, value, **style):
    cell = ws.cell(row=row, column=col, value=_clean(value))
    if isinstance(value, str):
        cell.data_type = "s"  # always text: a name that starts with "=" must never become a formula
    for k, v in style.items():
        setattr(cell, k, v)
    return cell


def to_xlsx(r: Report) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    head = PatternFill("solid", fgColor="1F3A5F")

    def new_sheet(name, title):
        ws = wb.create_sheet(_sheet_name(name, used))
        _put(ws, 1, 1, title, font=Font(bold=True, size=14))
        _put(ws, 2, 1, _sub_title(r), font=Font(italic=True, color="555555"))
        _put(ws, 3, 1, _made_at(), font=Font(italic=True, color="555555"))
        return ws

    if r.summary:
        ws = new_sheet("Summary", r.title)
        for i, (k, v) in enumerate(r.summary, start=5):
            _put(ws, i, 1, k, font=Font(bold=True))
            _put(ws, i, 2, v, alignment=Alignment(horizontal="right"))
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 18

    for s in r.sections:
        ws = new_sheet(s.title or r.title, f"{r.title}" + (f" - {s.title}" if s.title else ""))
        top = 5
        for c, col in enumerate(s.columns, start=1):
            _put(ws, top, c, col.label, font=Font(bold=True, color="FFFFFF"), fill=head,
                 alignment=Alignment(wrap_text=True, vertical="center"))
        for i, row in enumerate(s.rows, start=top + 1):
            for c, col in enumerate(s.columns, start=1):
                v = row.get(col.key)
                _put(ws, i, c, v, alignment=Alignment(horizontal="right") if isinstance(v, (int, float)) else None)
        for c, col in enumerate(s.columns, start=1):
            longest = max([len(col.label)] + [len(str(row.get(col.key) or "")) for row in s.rows[:300]])
            ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 8), 60)
        ws.freeze_panes = ws.cell(row=top + 1, column=1)
        if s.rows:
            ws.auto_filter.ref = f"A{top}:{get_column_letter(len(s.columns))}{top + len(s.rows)}"

    if not wb.sheetnames:
        wb.create_sheet("Report")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF -----------------------------------------------------------------------------------------------

_fonts: tuple[str, str] | None = None


def _pdf_fonts() -> tuple[str, str]:
    """(regular, bold) font names. Uses PDF_FONT_PATH / PDF_FONT_BOLD_PATH when set, else built-in Helvetica."""
    global _fonts
    if _fonts is None:
        _fonts = ("Helvetica", "Helvetica-Bold")
        if settings.pdf_font_path:
            pdfmetrics.registerFont(TTFont("ReportFont", settings.pdf_font_path))
            bold = "ReportFont"
            if settings.pdf_font_bold_path:
                pdfmetrics.registerFont(TTFont("ReportFont-Bold", settings.pdf_font_bold_path))
                bold = "ReportFont-Bold"
            _fonts = ("ReportFont", bold)
    return _fonts


def _text(v) -> str:
    return "" if v is None else escape(_clean(str(v)))


def to_pdf(r: Report) -> bytes:
    regular, bold = _pdf_fonts()
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=15, leading=19)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=3)
    small = ParagraphStyle("small", fontName=regular, fontSize=8, leading=10, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("cell", fontName=regular, fontSize=7.5, leading=9)
    cell_head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=colors.white)
    right = ParagraphStyle("right", parent=cell, alignment=2)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=r.title, author="Scan-to-Drive")
    width = landscape(A4)[0] - 24 * mm

    def footer(canvas, _doc):
        canvas.setFont(regular, 7)
        canvas.drawString(12 * mm, 7 * mm, f"Scan-to-Drive - {r.title} - {_made_at()}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Page {_doc.page}")

    story = [Paragraph(_text(r.title), h1), Paragraph(_text(_sub_title(r)), small), Spacer(1, 4 * mm)]
    if r.summary:
        data = [[Paragraph(f"<b>{_text(k)}</b>", cell), Paragraph(_text(v), right)] for k, v in r.summary]
        t = Table(data, colWidths=[75 * mm, 30 * mm], hAlign="LEFT")
        t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story += [t, Spacer(1, 3 * mm)]

    for s in r.sections:
        if s.title:
            story.append(Paragraph(_text(s.title), h2))
        cols = [c for c in s.columns if c.pdf]
        if not s.rows:
            story.append(Paragraph("No records in this period.", small))
            continue
        data = [[Paragraph(_text(c.label), cell_head) for c in cols]]
        for row in s.rows:
            data.append([Paragraph(_text(row.get(c.key)), right if isinstance(row.get(c.key), (int, float)) else cell)
                         for c in cols])
        weights = [min(max(len(c.label), max((len(str(row.get(c.key) or "")) for row in s.rows[:100]), default=0), 4), 40)
                   for c in cols]
        total = sum(weights)
        t = Table(data, colWidths=[width * w / total for w in weights], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
