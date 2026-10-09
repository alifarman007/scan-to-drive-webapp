"""Reports (7), Excel / PDF export, audit export and the purposes list."""
import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app.models import Alert, AlertStatus, AlertType, Passenger, Trip, TripStatus
from app.routers import admin_reports
from app.services import report_export
from tests.test_admin_trips import hdr, make_admin
from tests.test_end_trip import client, make_passenger  # noqa: F401  (client fixture)
from tests.test_trips import auth, make_car, make_driver

UTC = timezone.utc
# 2026-03: every test uses its own new car and driver, so data already in the development database never mixes in.
D1 = datetime(2026, 3, 10, 4, 0, tzinfo=UTC)  # 10:00 in Bangladesh
WINDOW = {"date_from": "2026-03-01", "date_to": "2026-03-31"}


def add_trip(db, car, driver, day=0, start_km=1000, end_km=1050, status=TripStatus.completed, passenger=None,
             hours=1, with_passenger=True, **kw):
    start = D1 + timedelta(days=day)
    if not with_passenger and passenger is None:
        kw.setdefault("purpose", "errand")  # the database wants a purpose for trips without a passenger
    if status in (TripStatus.cancelled, TripStatus.closed_by_admin):
        kw.setdefault("close_reason", "test reason")
    done = status in (TripStatus.completed, TripStatus.closed_by_admin) and end_km is not None
    t = Trip(
        vehicle_id=car.id, driver_id=driver.id, passenger_id=passenger.id if passenger else None,
        with_passenger=with_passenger if passenger is None else True, start_km=start_km,
        end_km=end_km if done else None, distance_km=end_km - start_km if done else None,
        start_place="Head Office", destination="Gazipur", start_time=start,
        end_time=start + timedelta(hours=hours) if done else None, status=status, **kw)
    db.add(t)
    db.flush()
    return t


def report(client, admin, name, **params):
    return client.get(f"/api/admin/reports/{name}", params=params, headers=hdr(admin))


def rows_of(body, section=0):
    return body["sections"][section]["rows"]


def summary_of(body):
    return {s["label"]: s["value"] for s in body["summary"]}


# ---- access ---------------------------------------------------------------------------------------

def test_report_list_and_access(client, db):
    admin, viewer = make_admin(db), make_admin(db, __import__("app.models", fromlist=["AdminRole"]).AdminRole.viewer)
    r = client.get("/api/admin/reports", headers=hdr(viewer))
    assert r.status_code == 200 and len(r.json()["reports"]) == 7
    assert report(client, viewer, "car_usage").status_code == 200
    assert client.get("/api/admin/reports/car_usage").status_code == 401
    driver = make_driver(db)
    assert client.get("/api/admin/reports/car_usage", headers=auth("driver", driver.id)).status_code == 403
    assert report(client, admin, "nope").json()["detail"]["code"] == "REPORT_NOT_FOUND"
    bad = report(client, admin, "car_usage", date_from="2026-03-10", date_to="2026-03-01")
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "BAD_DATE_RANGE"


def test_default_period_is_this_month(client, db):
    admin = make_admin(db)
    body = report(client, admin, "monthly_summary").json()
    today = datetime.now(UTC).astimezone(timezone(timedelta(hours=6))).date()
    assert body["period"] == {"from": today.replace(day=1).isoformat(), "to": today.isoformat()}


# ---- 1 car usage ----------------------------------------------------------------------------------

def test_car_usage(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    add_trip(db, car, driver, day=0, start_km=1000, end_km=1050, hours=2)
    add_trip(db, car, driver, day=0, start_km=1050, end_km=1070, hours=1)   # same day
    add_trip(db, car, driver, day=3, start_km=1070, end_km=1100, hours=3)
    add_trip(db, car, driver, day=5, start_km=1100, end_km=None, status=TripStatus.cancelled)
    body = report(client, admin, "car_usage", car_id=car.id, **WINDOW).json()
    (row,) = rows_of(body)
    assert row["car_code"] == car.car_code and row["trips"] == 3 and row["km"] == 100
    assert row["hours_used"] == 6.0 and row["days_used"] == 2
    assert row["idle_days"] == 31 - 2 and row["km_per_trip"] == 33.3


# ---- 2 drivers ------------------------------------------------------------------------------------

def test_driver_report(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db, "Karim Ali")
    t1 = add_trip(db, car, driver, day=0, start_km=0, end_km=40, hours=1)
    add_trip(db, car, driver, day=1, start_km=40, end_km=100, hours=3)
    add_trip(db, car, driver, day=2, start_km=100, end_km=None, status=TripStatus.cancelled)
    db.add(Alert(trip_id=t1.id, vehicle_id=car.id, type=AlertType.long_trip, message="x"))
    db.flush()
    (row,) = rows_of(report(client, admin, "drivers", driver_id=driver.id, **WINDOW).json())
    assert (row["name"], row["trips"], row["km"], row["cancelled"], row["alerts"]) == ("Karim Ali", 2, 100, 1, 1)
    assert row["avg_trip_minutes"] == 120


# ---- 3 passenger / department ---------------------------------------------------------------------

def test_passenger_and_department(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    a, b = make_passenger(db, "Nadia"), make_passenger(db, "Omar")
    a.department, b.department = "Sales", "Accounts"
    db.flush()
    add_trip(db, car, driver, day=0, start_km=0, end_km=30, passenger=a)
    add_trip(db, car, driver, day=1, start_km=30, end_km=50, passenger=a)
    add_trip(db, car, driver, day=2, start_km=50, end_km=100, passenger=b)
    add_trip(db, car, driver, day=3, start_km=100, end_km=110, is_visitor=True, visitor_name="Mr Lee",
             visitor_phone="01711111111", with_passenger=True)
    add_trip(db, car, driver, day=4, start_km=110, end_km=125, with_passenger=False)
    body = report(client, admin, "passenger_department", driver_id=driver.id, **WINDOW).json()
    dept = {r["department"]: r for r in rows_of(body, 0)}
    assert dept["Sales"]["trips"] == 2 and dept["Sales"]["km"] == 50 and dept["Accounts"]["km"] == 50
    assert dept["Visitor"]["km"] == 10 and dept["(no passenger)"]["km"] == 15
    assert dept["Sales"]["km_share_percent"] == round(100 * 50 / 125, 1)
    people = {r["name"]: r for r in rows_of(body, 1)}
    assert people["Nadia"]["trips"] == 2 and people["Nadia"]["employee_id"] == a.employee_id
    assert "Mr Lee (01711111111)" in people and len(people) == 3  # the trip without a passenger is not a person

    only = report(client, admin, "passenger_department", driver_id=driver.id, department="sales", **WINDOW).json()
    assert [r["department"] for r in rows_of(only, 0)] == ["Sales"]


# ---- 4 trip log -----------------------------------------------------------------------------------

def test_trip_log_has_everything_and_a_link(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    p = make_passenger(db)
    t = add_trip(db, car, driver, passenger=p, purpose="Meeting")
    add_trip(db, car, driver, day=1, start_km=1050, end_km=None, status=TripStatus.cancelled, close_reason="wrong car")
    body = report(client, admin, "trip_log", driver_id=driver.id, **WINDOW).json()
    first, second = rows_of(body)
    assert first["trip_no"] == t.trip_no and first["car_code"] == car.car_code and first["distance_km"] == 50
    assert first["passenger"] == p.name and first["purpose"] == "Meeting"
    assert first["start_time"] == "2026-03-10 10:00" and first["details_link"].endswith(f"/admin/trips/{t.id}")
    assert second["status"] == "cancelled" and second["note"] == "wrong car"
    assert summary_of(body)["Trips"] == 2


# ---- 5 exceptions ---------------------------------------------------------------------------------

def test_exceptions(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    add_trip(db, car, driver, day=0, start_km=0, end_km=10)                                # fine, not listed
    add_trip(db, car, driver, day=1, start_km=10, end_km=None, status=TripStatus.cancelled, close_reason="oops")
    add_trip(db, car, driver, day=2, start_km=10, end_km=20, status=TripStatus.closed_by_admin,
             close_reason="forgot to end", needs_review=True)
    add_trip(db, car, driver, day=3, start_km=20, end_km=30, with_passenger=False, purpose="errand")
    add_trip(db, car, driver, day=4, start_km=30, end_km=40, needs_review=True, approval_status="pending",
             start_no_scan_reason="phone broken")
    db.add(Alert(vehicle_id=car.id, type=AlertType.km_gap, message="gap of 60 km", status=AlertStatus.open,
                 created_at=D1 + timedelta(days=2)))
    db.add(Alert(vehicle_id=car.id, type=AlertType.admin_closed, message="not listed twice",
                 created_at=D1 + timedelta(days=2)))
    db.flush()
    body = report(client, admin, "exceptions", car_id=car.id, **WINDOW).json()
    types = [r["type"] for r in rows_of(body)]
    assert sorted(types) == sorted(["Cancelled trip", "Closed by admin", "No passenger",
                                    "Needs review (passenger could not scan)", "Km gap"])
    gap = next(r for r in rows_of(body) if r["type"] == "Km gap")
    assert gap["details"] == "gap of 60 km" and gap["status"] == "open" and gap["trip_no"] == ""
    review = next(r for r in rows_of(body) if r["type"].startswith("Needs review"))
    assert "phone broken" in review["details"] and review["status"] == "pending"
    assert summary_of(body)["Exceptions"] == 5


# ---- 6 km continuity ------------------------------------------------------------------------------

def test_km_continuity_finds_gaps_and_uses_trip_before_the_period(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    add_trip(db, car, driver, day=-20, start_km=500, end_km=600)       # February, before the window
    add_trip(db, car, driver, day=0, start_km=600, end_km=650)          # OK, compared with the February trip
    add_trip(db, car, driver, day=1, start_km=700, end_km=720)          # 50 km nobody wrote down
    add_trip(db, car, driver, day=2, start_km=700, end_km=710)          # odometer went back 20
    add_trip(db, car, driver, day=3, start_km=710, end_km=None, status=TripStatus.cancelled)  # ignored
    body = report(client, admin, "km_continuity", car_id=car.id, **WINDOW).json()
    results = [r["result"] for r in rows_of(body)]
    assert results == ["OK", "Unrecorded use: 50 km", "Odometer went back by 20 km"]
    assert [r["gap_km"] for r in rows_of(body)] == [0, 50, -20]
    assert summary_of(body) == {"Trips checked": 3, "Trips with a gap": 2}
    first = rows_of(report(client, admin, "km_continuity", car_id=car.id, date_from="2026-02-01",
                           date_to="2026-03-31").json())[0]
    assert first["result"] == "First record for this car"


# ---- 7 monthly summary ----------------------------------------------------------------------------

def test_monthly_summary(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    p = make_passenger(db)
    p.department = "HR"
    db.flush()
    add_trip(db, car, driver, day=0, start_km=0, end_km=40, passenger=p)
    add_trip(db, car, driver, day=1, start_km=40, end_km=60, status=TripStatus.closed_by_admin, with_passenger=False)
    add_trip(db, car, driver, day=2, start_km=60, end_km=None, status=TripStatus.cancelled)
    body = report(client, admin, "monthly_summary", car_id=car.id, **WINDOW).json()
    s = summary_of(body)
    assert s["Trips (not cancelled)"] == 2 and s["Completed"] == 1 and s["Closed by admin"] == 1
    assert s["Cancelled"] == 1 and s["Total km"] == 60 and s["Average km per trip"] == 30.0
    assert s["Trips without a passenger"] == 1
    assert rows_of(body, 0) == [{"car_code": car.car_code, "trips": 2, "km": 60}]
    assert {r["department"] for r in rows_of(body, 1)} == {"HR", "(no passenger)"}


# ---- export ---------------------------------------------------------------------------------------

def test_excel_export(client, db):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db, "=HYPERLINK(1)")
    add_trip(db, car, driver, is_visitor=True, visitor_name="=1+1", visitor_phone="017")
    r = report(client, admin, "trip_log", format="xlsx", driver_id=driver.id, **WINDOW)
    assert r.status_code == 200 and r.headers["content-type"] == report_export.XLSX_TYPE
    assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Summary"
    ws = wb["Trip log"]
    assert ws["A1"].value == "Trip log" and "Period: 2026-03-01 to 2026-03-31" in ws["A2"].value
    header = [c.value for c in ws[5]]
    assert header[0] == "Trip" and "Km" in header
    row = {h: c.value for h, c in zip(header, ws[6])}
    assert row["Driver"] == "=HYPERLINK(1)" and ws.cell(6, header.index("Driver") + 1).data_type == "s"  # text, never a formula
    assert row["Passenger"] == "=1+1 (visitor)"
    assert row["Km"] == 50 and isinstance(row["Km"], int)


def test_two_section_excel(client, db):
    admin = make_admin(db)
    r = report(client, admin, "passenger_department", format="xlsx", **WINDOW)
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Summary", "By department", "By passenger"]


@pytest.mark.parametrize("name", ["car_usage", "drivers", "passenger_department", "trip_log", "exceptions",
                                  "km_continuity", "monthly_summary"])
def test_every_report_makes_both_files(client, db, name):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    add_trip(db, car, driver, passenger=make_passenger(db))
    x = report(client, admin, name, format="xlsx", car_id=car.id, **WINDOW)
    p = report(client, admin, name, format="pdf", car_id=car.id, **WINDOW)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert p.status_code == 200 and p.content[:5] == b"%PDF-" and p.headers["content-type"] == "application/pdf"
    assert ".pdf" in p.headers["content-disposition"]


def test_pdf_with_a_custom_font(client, db, monkeypatch):
    import os
    font = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    if not os.path.exists(font):
        pytest.skip("DejaVu font not installed")
    monkeypatch.setattr(report_export.settings, "pdf_font_path", font)
    monkeypatch.setattr(report_export, "_fonts", None)
    admin, car, driver = make_admin(db), make_car(db), make_driver(db, "Rahim é <b>&")
    add_trip(db, car, driver)
    r = report(client, admin, "drivers", format="pdf", driver_id=driver.id, **WINDOW)
    assert r.status_code == 200 and b"DejaVuSans" in r.content
    monkeypatch.setattr(report_export, "_fonts", None)


def test_export_too_big(client, db, monkeypatch):
    admin, car, driver = make_admin(db), make_car(db), make_driver(db)
    add_trip(db, car, driver)
    add_trip(db, car, driver, day=1, start_km=1050, end_km=1060)
    monkeypatch.setattr(admin_reports, "MAX_ROWS", 1)
    r = report(client, admin, "trip_log", format="xlsx", car_id=car.id, **WINDOW)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "TOO_MANY_ROWS"
    assert report(client, admin, "trip_log", format="json", car_id=car.id, **WINDOW).status_code == 200


# ---- audit export ---------------------------------------------------------------------------------

def test_audit_export_admin_only(client, db):
    from app.models import AdminRole
    admin, viewer = make_admin(db), make_admin(db, AdminRole.viewer)
    d = make_driver(db)
    wrong = client.post("/api/auth/driver/login", json={"employee_id": d.employee_id, "pin": "9999"})
    assert wrong.status_code == 401
    today = datetime.now(UTC).date().isoformat()
    r = client.get("/api/admin/audit/export", params={"actor": d.employee_id, "date_from": today},
                   headers=hdr(admin))
    assert r.status_code == 200 and r.content[:2] == b"PK"
    ws = load_workbook(io.BytesIO(r.content))["Audit log"]
    assert ws["A5"].value == "Time" and ws["B6"].value and ws["C6"].value == f"driver:{d.employee_id}"
    pdf = client.get("/api/admin/audit/export", params={"format": "pdf", "actor": d.employee_id}, headers=hdr(admin))
    assert pdf.content[:5] == b"%PDF-"
    assert client.get("/api/admin/audit/export", headers=hdr(viewer)).status_code == 403


# ---- purposes list --------------------------------------------------------------------------------

def test_purposes_list(client, db):
    from app.models import AdminRole
    admin, viewer, driver = make_admin(db), make_admin(db, AdminRole.viewer), make_driver(db)
    assert client.get("/api/purposes").status_code == 401
    assert client.get("/api/purposes", headers=auth("driver", driver.id)).json()["purposes"]  # a default list at least
    r = client.put("/api/admin/settings/purposes", headers=hdr(admin),
                   json={"purposes": ["  Client   visit ", "Bank", "client visit", "Airport"]})
    assert r.status_code == 200 and r.json()["purposes"] == ["Client visit", "Bank", "Airport"]
    assert client.get("/api/purposes", headers=auth("driver", driver.id)).json()["purposes"] == ["Client visit", "Bank", "Airport"]
    assert client.get("/api/admin/settings", headers=hdr(viewer)).json()["purposes"] == ["Client visit", "Bank", "Airport"]
    assert client.put("/api/admin/settings/purposes", headers=hdr(viewer), json={"purposes": ["x"]}).status_code == 403
    assert client.put("/api/admin/settings/purposes", headers=hdr(admin), json={"purposes": []}).status_code == 422
    too_long = client.put("/api/admin/settings/purposes", headers=hdr(admin), json={"purposes": ["x" * 61]})
    assert too_long.status_code == 422 and too_long.json()["detail"]["code"] == "BAD_PURPOSE"
    blank = client.put("/api/admin/settings/purposes", headers=hdr(admin), json={"purposes": ["  "]})
    assert blank.status_code == 422
